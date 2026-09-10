import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x00, 0x7A, 0xFF)
GREEN = RGBColor(0x0E, 0x9F, 0x6E)
RED = RGBColor(0xC0, 0x39, 0x2B)
DARK = RGBColor(0x1F, 0x29, 0x37)
GREY = RGBColor(0x6B, 0x72, 0x80)

doc = Document()

# ---- base styles ----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK

def set_h(style_name, size, color=ACCENT, bold=True):
    s = doc.styles[style_name]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.color.rgb = color

set_h("Heading 1", 16)
set_h("Heading 2", 13)
set_h("Heading 3", 11.5, DARK)

def para(text, bold=False, italic=False, color=None, size=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    return p

def bullet(text, bold_prefix=None, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(2)
    return p

def numbered(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Number")
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(2)
    return p

def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def make_table(headers, rows, widths=None, header_fill="007AFF"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9.5)
        shade_cell(hdr[i], header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9)
    if widths:
        for i, w in enumerate(widths):
            for r in t.rows:
                r.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

# ================= COVER =================
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ADHERA (MediTrack)"); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = ACCENT
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Project Progress Report"); r.font.size = Pt(18); r.font.color.rgb = DARK
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(f"Updated: {datetime.date(2026, 9, 8).strftime('%B %d, %Y')}   |   Repository: github.com/yoofisey/MediTrack   |   Domain: www.useadhera.com")
r.font.size = Pt(10); r.font.color.rgb = GREY
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Release-readiness edition: security audit, secret hygiene, error monitoring, payments (Paystack + Paddle), branded native assets, and a staged roadmap to store release")
r.italic = True; r.font.size = Pt(10); r.font.color.rgb = GREY
doc.add_paragraph()

# ================= TOC =================
doc.add_heading("Table of Contents", 1)
for i, item in enumerate([
    "1. Executive Summary",
    "2. Project Overview & Tech Stack",
    "3. Feature Summary",
    "4. What Was Completed (Sept 2026 cycles)",
    "5. Readiness & Security Audit - Results",
    "6. Payments Status (Paystack + Paddle)",
    "7. Release Roadmap to Store Distribution",
    "8. Open Items & Owner",
    "9. Versioning & Signing Reference",
], 1):
    p = doc.add_paragraph(); r = p.add_run(item); r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(2)
doc.add_page_break()

# ================= 1. EXEC SUMMARY =================
doc.add_heading("1. Executive Summary", 1)
para("Adhera (formerly MediTrack) is a mobile-first medication management and health-tracking platform for the West African market (Ghana, Nigeria, with South Africa and Kenya planned). It is built on Next.js (web/PWA), Capacitor (native iOS/Android shells that load the hosted web app), and Supabase (auth + Postgres with Row Level Security). Freemium billing runs through Paystack, with Paddle as the Merchant of Record for every other market.", space_after=4)
para("The application reached feature-complete status and this project is now in release preparation. Completed this cycle: a full launch-readiness and security audit, server-side secret hygiene fix, production error monitoring (Sentry), branded native launcher icons, documented Codemagic signing/CI, country-aware Paystack billing (Ghana live; NG/ZA/KE activate via env), a new Paddle Merchant-of-Record checkout for international payments, and the Content-Security-Policy update required by the Paddle overlay. The production web build compiles cleanly, the Android release build succeeds, and Sentry + all five Paddle environment variables are now configured on Vercel Production. Remaining work is concentrated in external/account steps (store assets, Apple Developer membership, Play rollout, domain/RLS/push verification).", space_after=4)

# ================= 2. OVERVIEW =================
doc.add_heading("2. Project Overview & Tech Stack", 1)
para("2.1 Product Vision", bold=True, color=ACCENT)
para("A comprehensive medication adherence app with a freemium model (Free / Pro GHS 15/mo / Family GHS 28/mo at Ghana pricing; international equivalents via Paddle), supporting individual medication tracking, family/caregiver management, community features, health journaling, symptom tracking, gamification (badges and challenges), and premium analytics.")
para("2.2 Tech Stack", bold=True, color=ACCENT)
make_table(
    ["Layer", "Technology"],
    [
        ["Frontend", "Next.js 16 (App Router), React 19, Turbopack"],
        ["Mobile shells", "Capacitor 8 (iOS + Android) loading hosted web app (webDir: out)"],
        ["Backend / DB", "Supabase (Postgres + RLS), Next.js API routes on Vercel"],
        ["Auth", "Supabase Auth (email/OTP; implicit flow storage key mt_sb_session)"],
        ["Payments", "Paystack (GHS/NGN/ZAR/KES) with signed webhook verification + Paddle (MoR, all other countries)"],
        ["Push", "Firebase Cloud Messaging (server-side) + Web Push VAPID"],
        ["CI/CD", "Codemagic (Android/iOS release), Vercel (web + cron)"],
        ["Error monitoring", "Sentry (@sentry/nextjs 10) - client/server/edge + tunnel"],
        ["Branding", "Adhera; bundle id com.useadhera.app"],
    ],
    widths=[1.6, 4.9],
)
para("2.3 Target Platforms", bold=True, color=ACCENT)
para("Web (PWA) - www.useadhera.com (live on Vercel); Android (Capacitor + Codemagic); iOS (Capacitor - pending Apple Developer account).")

# ================= 3. FEATURES =================
doc.add_heading("3. Feature Summary", 1)
make_table(
    ["Area", "Capabilities"],
    [
        ["Medication tracking", "List, dose scheduling, daily view, adherence analytics, visit history (month-grouped + 3-recent teaser)"],
        ["Filters", "Search / status / date-range filtering on visit history"],
        ["Reminders", "Smart alarms/reminders, condition & schedule editors, push toggle"],
        ["Health data", "Vitals, journaling, symptom tracking, medication cost tracking, prescription columns"],
        ["Family / caregiver", "Family dashboard, caregiver home, family vitals, dedicated family tier"],
        ["Community", "Community feed, gamification (badges, challenges)"],
        ["Auth & profile", "Sign-up/login/OTP, reset password, profile & avatar, medical ID"],
        ["Payments", "Freemium tiers; Paystack (GH live, NG/ZA/KE env-ready) + Paddle international; signed webhooks"],
        ["Premium UX", "Glassmorphism, dark mode, premium splash, i18n wiring"],
    ],
    widths=[1.6, 4.9],
)

# ================= 4. COMPLETED =================
doc.add_heading("4. What Was Completed (September 2026)", 1)
para("Two work streams: (1) launch-readiness/security hardening and native asset + CI preparation; (2) completing the international payments story so non-Ghana customers can pay. All builds verify green (tsc clean, next build success).", space_after=4)

doc.add_heading("4.1 Visit History + Filters (committed: 7f5bb04, e500d4f)", 2)
para("Shipped a Visit History feature: a three-recent Reports teaser plus a full month-grouped visit page, and added search / status / date-range filters.", space_after=4)

doc.add_heading("4.2 Launch-Readiness Audit + Code Fixes (4bd74d8, 73db3d4, 105075a)", 2)
bullet("", bold_prefix="Versions unified to 1.0.0: ")
para("package.json, Android versionName '1.0.0' (versionCode 1), iOS MARKETING_VERSION 1.0.0.")
bullet("", bold_prefix="iOS privacy strings: ")
para("NSCameraUsageDescription (QR scanning for doctor sharing) and NSPhotoLibraryUsageDescription (avatar/profile pictures).")
bullet("", bold_prefix="Dead code removed: ")
para("MeTab.jsx and EnterprisePaymentView.jsx; About version now reads live from @capacitor/app.")
bullet("", bold_prefix="Cap sync + lint fixes: ")
para("synced Capacitor config including LocalNotifications icon/color; fixed ProfileTab import and lazy Date.now() initializer; removed unused imports/medCount prop (eslint clean).")

doc.add_heading("4.3 Android Signing Keystore (regenerated 2026-09-03)", 2)
para("A fresh RSA-2048 keystore (alias 'adhera', 10,000-day validity) was generated and verified: assembleRelease BUILD SUCCESSFUL and apksigner confirms SHA-256 81:85:E5:43:BE:77:B9:23:2A:9E:F2:61:82:73:EA:07:0C:1D:54:A8:88:88:7B:99:C7:47:5E:CB:F6:40:35:52. Backed up at C:\\Users\\franc\\Documents\\Adhera-Signing\\ (adhera-release.jks, adhera-cert.pem, keystore.b64). Keystore.properties is gitignored.")

doc.add_heading("4.4 Branded Native Launcher Icons (9f43ddb)", 2)
para("Generated adaptive Android launcher icons from public/icon.svg: a white 'A + medical cross' foreground glyph across all five densities, a gradient drawable background (#007AFF -> #5856D6), plus legacy PNGs. Verified all densities non-empty, glyph centered, and assembleRelease still builds.")

doc.add_heading("4.5 Codemagic Signing Docs (310fc00)", 2)
para("docs/mobile-ci.md updated to reference the new keystore backup location, record the alias + certificate SHA-256, and describe the android_keystore / google_services / appstore_connect secret groups.")

doc.add_heading("4.6 Security Hardening - Paystack Secret Split (714b795)", 2)
para("Audit found PAYSTACK_SECRET_KEY was defined in lib/payments.js, a module imported by 'use client' components (Modals.jsx, PricingPage.jsx). Although only called server-side, the module-level constant risked leaking into the browser bundle. Fix: moved secrets + getPaystackSecret() into a dedicated server-only module (lib/payments-server.ts); the shared module now exposes only NEXT_PUBLIC_* config. Verified with lint + tsc.")

doc.add_heading("4.7 Error Monitoring - Sentry (52deab1, DSN configured on Vercel)", 2)
para("Installed @sentry/nextjs 10.73 and wired client, server, and edge runtimes: sentry.client/server/edge.config.ts, instrumentation.ts (register + onRequestError) and instrumentation-client.ts (onRouterTransitionStart). Added app/error.tsx as the root error boundary. Wrapped next.config.ts with withSentryConfig using a tunnel route (/sentry-tunnel) so browser events travel through the app's own origin - the strict health-app Content-Security-Policy did NOT need relaxing. NEXT_PUBLIC_SENTRY_DSN + SENTRY_DSN are set on Vercel Production/Preview/Development (Config type). Remaining (optional): SENTRY_ORG / SENTRY_PROJECT / SENTRY_AUTH_TOKEN for source-map upload and issue alerting.")

doc.add_heading("4.8 Country-Aware Paystack Gating (8f0f487)", 2)
para("Paystack plumbing is now environment-driven per market. Ghana plans (PLN_w5rq3bkd5uh5mqj pro, PLN_h9mlqfmujuh74c9 family) remain hard-coded and live; NG/ZA/KE become billable automatically once their public key, secret key, and plan codes are set in Vercel - no code changes. Server-side: per-country secrets map (PAYSTACK_SECRET_KEY[_NG/_ZA/_KE]), currency map (GHS/NGN/ZAR/KES), and PLAN_MIN_AMOUNTS to fence price edits. Webhook derives the secret from metadata.country; the verify endpoint accepts { reference, country }. .env.example documents every per-market variable.")

doc.add_heading("4.9 Paddle Merchant-of-Record Checkout (5adbce8, a3b2a56)", 2)
para("For every country outside the Paystack markets (GH/NG/ZA/KE), checkout now routes to Paddle, a Merchant of Record that handles sales tax/VAT, chargebacks, and pays out in USD. getPaymentsConfig() returns a 'paddle' gateway with client token + price IDs. Client: Paddle.js v2 loaded on demand and Paddle.Checkout.open() with customData { plan, user_id } so the webhook can map the subscription to the Supabase user. Server: /api/paddle/webhook verifies the Paddle-Signature header (HMAC-SHA256 over ts:rawBody, 5-minute freshness, timingSafeEqual), provisions from subscription.created/activated/updated + transaction.completed (writes payment_references + profiles.plan/paid_at), and downgrades to free on cancel/past-due/pause. next.config.ts CSP was updated to allow cdn.paddle.com (script/style/img), *.paddle.com (connect), and the buy.paddle.com / sandbox-buy.paddle.com checkout iframes.")

# ================= 5. AUDIT =================
doc.add_heading("5. Readiness & Security Audit - Results", 1)
para("A five-stream audit (auth/OAuth, API routes/authz, Supabase RLS, client XSS/secrets, dependencies/headers) was run and the top claims were verified hands-on. Findings below are the confirmed status.")

doc.add_heading("5.1 Confirmed & Fixed / Verified Secure", 2)
make_table(
    ["Finding", "Status", "Detail"],
    [
        ["Paystack secret in client bundle", "FIXED", "Split to server-only module (714b795)"],
        ["Cron endpoints unauthenticated", "VERIFIED SAFE", "All cron routes guard on verifyCronAuth (fail-open only if CRON_SECRET unset, i.e. dev)"],
        ["Paystack webhook trust", "VERIFIED SAFE", "HMAC-SHA512 signature + timingSafeEqual + plan allowlist + payment_references dedup"],
        ["Paddle webhook trust", "VERIFIED SAFE", "Paddle-Signature HMAC-SHA256 (ts:body) + 5-min freshness + timingSafeEqual (a3b2a56)"],
        [".env.vercel secret in git", "VERIFIED SAFE", ".env.vercel is gitignored; only .env.example is tracked"],
        ["PDF preview XSS", "LOW RISK", "dangerouslySetInnerHTML fed through escapeHtml() on all user fields"],
    ],
    widths=[2.3, 1.4, 2.8],
)

doc.add_heading("5.2 Residual / Actionable (config, not code)", 2)
bullet("", bold_prefix="Confirm CRON_SECRET is set in Vercel production. ")
para("Until set, cron endpoints are unauth-accessible in dev. This is a config check, not a code change.")
bullet("", bold_prefix="Paddle checkout domain approval. ")
para("Live Paddle accounts require www.useadhera.com to be approved under Paddle > Checkout > Website approval before the overlay will open.")

# ================= 6. PAYMENTS =================
doc.add_heading("6. Payments Status (Paystack + Paddle)", 1)
para("Decision (Sept 2026): Ghana launches on Paystack; NG/ZA/KE become billable later by setting per-market env vars; every other country pays through Paddle (MoR). This makes the app usable internationally without opening multiple merchant accounts.", space_after=4)

make_table(
    ["Gateway", "Countries", "Status", "How to activate"],
    [
        ["Paystack", "GH", "LIVE (plans hard-coded)", "None - plans PLN_... active; international-card toggle optional (dashboard)"],
        ["Paystack", "NG / ZA / KE", "Staged (coming soon)", "Set PAYSTACK_SECRET_KEY_x + NEXT_PUBLIC_PAYSTACK_KEY_x + plan codes on Vercel"],
        ["Paddle", "All other countries", "CONFIGURED on Vercel", "5 env vars live; needs checkout-domain approval + live webhook delivery test"],
    ],
    widths=[1.1, 1.0, 1.5, 2.9],
)
bullet("", bold_prefix="Vercel env vars set (Production/Preview/Development): ")
para("NEXT_PUBLIC_PAYSTACK_KEY + PAYSTACK_SECRET_KEY (GH, live); NEXT_PUBLIC_SENTRY_DSN + SENTRY_DSN; NEXT_PUBLIC_PADDLE_CLIENT_TOKEN, NEXT_PUBLIC_PADDLE_ENV=production, NEXT_PUBLIC_PADDLE_PRICE_PRO (pri_01m1zce9p0zrjz6gway2y5gmtr), NEXT_PUBLIC_PADDLE_PRICE_FAMILY (pri_01m1zcc0xr1tkepzwebpezpd59), PADDLE_WEBHOOK_SECRET (Secret type).")

# ================= 7. ROADMAP =================
doc.add_heading("7. Release Roadmap to Store Distribution", 1)
para("The web build is release-ready (Paystack + Paddle both configured). The path to Play Store / App Store distribution is staged below. Each stage lists exactly what must happen and the expected outcome before proceeding.", space_after=6)

def stage(title, owner, gate, steps):
    doc.add_heading(title, 2)
    p = doc.add_paragraph()
    r = p.add_run("Owner: "); r.bold = True; r.font.color.rgb = ACCENT
    p.add_run(owner)
    p = doc.add_paragraph()
    r = p.add_run("Exit gate: "); r.bold = True; r.font.color.rgb = GREEN
    p.add_run(gate)
    for s in steps:
        bullet("", bold_prefix=(" - " if False else "") + s)

stage("Stage 1 - Accounts, Secrets & Payments Verification (mostly complete)",
      "You",
      "All CI secrets, accounts, and production environment variables are configured AND verified working end-to-end (payments webhooks, Sentry, push).",
      [
        "Verify CRON_SECRET is set on Vercel Production (cron endpoints) - last config item.",
        "Globally test both payment paths: a GH Paystack card/mobile-money payment and an international Paddle checkout; confirm webhook logs in Paddle > Developer tools > Notifications.",
        "Approve www.useadhera.com for Paddle checkout (Checkout > Website approval) so the live overlay opens.",
        "Enable the Paystack 'International Payments' toggle (Settings > Business Settings) to accept foreign cards on the GH merchant, if desired.",
        "Paste the GOOGLE_SERVICES_JSON value into the Codemagic google_services secret group (value was computed: project adhera-ce56d, storage adhera-ce56d.firebasestorage.app).",
      ])

stage("Stage 2 - Store Visual Assets & Listing (external)",
      "You (with design)",
      "Play Console and App Store Connect are ready to accept a build (all store metadata + graphics uploaded, no missing-asset errors).",
      [
        "Capture screenshots on a real device for each supported form factor (phone; Android tablets are optional).",
        "Produce the 512x512 / 1024x1024 icon, feature graphic (1024x500), and promo/video frames.",
        "Upload store metadata from store-listing.md (short + full description, keywords, categories).",
        "Set up privacy policy URL and data-safety / nutrition labels.",
      ])

stage("Stage 3 - Android: CI Build & Internal Test (Codemagic)",
      "You",
      "A Codemagic android-release build succeeds and the produced AAB/APK is verified signed with the Adhera key (SHA-256 81:85:E5:43:...:52) and installs on a test device.",
      [
        "Populate the android_keystore group (KEYSTORE_BASE64, KEYSTORE_PASSWORD, KEY_ALIAS=adhera, KEY_PASSWORD) and the google_services group in Codemagic.",
        "Push to main / trigger the android-release workflow; download the AAB from artifacts.",
        "Verify signature matches the local keystore's SHA-256 and install on a device.",
      ])

stage("Stage 4 - Android: Play Console Rollout",
      "You",
      "Internal testing approved; build promoted through Closed -> Open -> Production, or a service account automation enabled.",
      [
        "Create the Play Console app with package com.useadhera.app.",
        "Upload the .aab under Testing -> Internal testing, add testers, review release notes.",
        "Promote through Closed / Open / Production tracks.",
        "(Optional) For automated uploads, create a Google Play service account and add GCLOUD_SERVICE_ACCOUNT_CREDENTIALS, then uncomment the google_play block in codemagic.yaml.",
      ])

stage("Stage 5 - iOS: Apple Developer Membership & Signing (external)",
      "You",
      "Apple Developer membership active; bundle ID com.useadhera.app registered; app record created in App Store Connect; App Store Connect API key generated.",
      [
        "Enroll in the Apple Developer Program ($99/yr).",
        "Register the bundle ID com.useadhera.app in the Developer portal.",
        "Generate an App Store Connect API key (App Manager role); record Key ID and Issuer ID.",
        "Populate the appstore_connect Codemagic group (APP_STORE_CONNECT_PRIVATE_KEY, KEY_ID, ISSUER_ID).",
      ])

stage("Stage 6 - iOS: Build, Push Cert, TestFlight",
      "You",
      "Codemagic ios-release (or manual Xcode archive) produces an archive; TestFlight build uploads successfully; Firebase APNs Auth Key attached for push.",
      [
        "Attach an APNs Auth Key (Firebase Console -> Cloud Messaging) to enable iOS push through FCM.",
        "Upload the iOS build to TestFlight (auto via codemagic app_store_connect block, or manually via Transporter).",
        "Confirm aps-environment is 'production' in the uploaded build.",
      ])

stage("Stage 7 - Domain & Deep-Link / RLS / Push Verification",
      "You",
      "All external verifications pass; app links and push resolve in production.",
      [
        "Verify domain ownership of useadhera.com (and www) - required for assetlinks.json and apple-app-site-association.",
        "Ensure assetlinks.json (Android app links) and apple-app-site-association (iOS) are served at the domain root.",
        "Apply pending Supabase RLS migrations to production and confirm policies are enforced.",
        "Test FCM push on a real Android and iOS device (if iOS available).",
      ])

# ================= 8. OPEN ITEMS =================
doc.add_heading("8. Open Items & Owner", 1)
make_table(
    ["Item", "Owner", "Type"],
    [
        ["Verify CRON_SECRET set on Vercel prod", "You", "Config/verification"],
        ["Globally test Paystack (GH) + Paddle (international) live payments", "You", "Verification"],
        ["Approve Paddle checkout domain (www.useadhera.com)", "You", "External"],
        ["Enable Paystack International Payments toggle (optional)", "You", "Config"],
        ["Paste GOOGLE_SERVICES_JSON into Codemagic google_services group", "You", "Config"],
        ["Optional: Sentry SENTRY_ORG / SENTRY_PROJECT / SENTRY_AUTH_TOKEN", "You", "Config (optional)"],
        ["Store visual assets (screenshots, feature graphic, 1024 icon)", "You / design", "External"],
        ["Apple Developer membership + bundle ID + App Store Connect", "You", "External, paid"],
        ["iOS push: APNs Auth Key in Firebase", "You", "External"],
        ["Domain/RLS/push production verification", "You", "External"],
        ["Google Play service account automation (optional)", "You", "Config"],
    ],
    widths=[3.0, 1.6, 1.9],
)

# ================= 9. VERSIONING =================
doc.add_heading("9. Versioning & Signing Reference", 1)
make_table(
    ["Fact", "Value"],
    [
        ["App version", "1.0.0 (web / Android / iOS aligned)"],
        ["Android applicationId", "com.useadhera.app"],
        ["iOS PRODUCT_BUNDLE_IDENTIFIER", "com.useadhera.app"],
        ["Android signing alias", "adhera"],
        ["Certificate SHA-256", "81:85:E5:43:BE:77:B9:23:2A:9E:F2:61:82:73:EA:07:0C:1D:54:A8:88:88:7B:99:C7:47:5E:CB:F6:40:35:52"],
        ["Keystore backup", "C:\\Users\\franc\\Documents\\Adhera-Signing\\"],
        ["Paystack GH plan codes", "pro PLN_w5rq3bkd5uh5mqj | family PLN_h9mlqfmujuh74c9"],
        ["Paddle price IDs", "pro pri_01m1zce9p0zrjz6gway2y5gmtr | family pri_01m1zcc0xr1tkepzwebpezpd59"],
        ["Web build output", "next build - green; static pages + server API routes on Vercel"],
    ],
    widths=[2.3, 4.2],
)

para("")
para("End of report.", italic=True, color=GREY)
para(f"Generated automatically: {datetime.date(2026, 9, 8).strftime('%B %d, %Y')}", color=GREY, size=9)

out = "Adhera_Progress_Report.docx"
doc.save(out)
print("Wrote", out)